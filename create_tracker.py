#!/usr/bin/env python3
"""Create the Excel action tracker template (with realistic SAMPLE data).

Sample dates are generated relative to today so the demo always shows a mix of
overdue, blocked and upcoming items. Replace the sample rows with real data.
"""
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).parent / "data" / "action_tracker.xlsx"
FONT = "Arial"

HEADERS = [
    "ID", "Initiative", "Action Item", "Owner", "Function", "Priority", "Status",
    "Start Date", "Due Date", "% Complete", "Last Update", "Remarks",
    "Days to Due", "RAG",
]
INPUT_COLS = 12  # A..L are typed in; M..N are formulas

FUNCTIONS = ["Business", "Operations", "Finance", "Risk", "Compliance", "Technology", "Infrastructure"]
STATUSES = ["Not Started", "In Progress", "Blocked", "Completed"]
PRIORITIES = ["High", "Medium", "Low"]

# (initiative, action, owner, function, priority, status, start_off, due_off, pct, update_off, remarks)
SAMPLE = [
    ("Client Onboarding Digitisation", "Finalise process map for e-onboarding", "Aarav Mehta", "Operations", "High", "Completed", -40, -10, 1.0, -10, ""),
    ("Client Onboarding Digitisation", "Sign off UAT test scripts", "Neha Iyer", "Technology", "High", "In Progress", -20, -4, 0.7, -2, "Two test cases pending"),
    ("Client Onboarding Digitisation", "Compliance review of digital consent flow", "Rohan Das", "Compliance", "High", "Blocked", -25, -6, 0.4, -9, "Awaiting legal opinion"),
    ("Regulatory Reporting Automation", "Reconcile data fields to regulatory template", "Priya Nair", "Finance", "High", "In Progress", -15, 3, 0.6, -1, ""),
    ("Regulatory Reporting Automation", "Automate monthly return generation", "Kunal Shah", "Technology", "Medium", "In Progress", -10, 12, 0.3, -3, ""),
    ("Regulatory Reporting Automation", "Document control framework for reports", "Sana Qureshi", "Risk", "Medium", "Not Started", 0, 20, 0.0, -5, ""),
    ("Cost Optimisation FY27", "Baseline cost per transaction by product", "Meera Pillai", "Finance", "High", "Completed", -30, -5, 1.0, -5, ""),
    ("Cost Optimisation FY27", "Identify vendor contract savings", "Vikram Rao", "Operations", "Medium", "In Progress", -12, -2, 0.5, -4, "Two contracts still under review"),
    ("Cost Optimisation FY27", "Present savings options to steering committee", "Aarav Mehta", "Business", "High", "Not Started", 2, 6, 0.0, -2, ""),
    ("Vendor Consolidation", "Shortlist vendors via RFP scoring", "Vikram Rao", "Operations", "Medium", "In Progress", -18, 9, 0.55, -2, ""),
    ("Vendor Consolidation", "Security assessment of shortlisted vendors", "Ishaan Verma", "Infrastructure", "High", "Blocked", -8, 5, 0.2, -6, "Vendor has not shared SOC 2 report"),
    ("Audit Findings Remediation", "Close access-review findings", "Sana Qureshi", "Risk", "High", "In Progress", -35, -14, 0.8, -3, "Evidence for 3 findings outstanding"),
    ("Audit Findings Remediation", "Update policy documents", "Rohan Das", "Compliance", "Medium", "Completed", -28, -7, 1.0, -7, ""),
    ("Audit Findings Remediation", "Compile evidence pack for internal audit", "Neha Iyer", "Technology", "Medium", "In Progress", -5, 4, 0.4, -1, ""),
    ("KYC Refresh Programme", "Prioritise high-risk client segment list", "Priya Nair", "Compliance", "High", "Completed", -22, -3, 1.0, -3, ""),
    ("KYC Refresh Programme", "Set up outreach tracker for client documents", "Kavya Reddy", "Business", "Medium", "In Progress", -14, 7, 0.65, -1, ""),
    ("KYC Refresh Programme", "Define escalation protocol for non-responsive clients", "Kavya Reddy", "Business", "Low", "Not Started", 3, 25, 0.0, -4, ""),
    ("Productivity Dashboard Rollout", "Define KPI dictionary", "Meera Pillai", "Finance", "High", "Completed", -20, -8, 1.0, -8, ""),
    ("Productivity Dashboard Rollout", "Build prototype dashboard", "Ishaan Verma", "Technology", "Medium", "In Progress", -7, 1, 0.45, -1, ""),
    ("Productivity Dashboard Rollout", "Run stakeholder training sessions", "Aarav Mehta", "Business", "Low", "Not Started", 5, 30, 0.0, -2, ""),
    ("Productivity Dashboard Rollout", "Automate reminder emails for open actions", "Kunal Shah", "Technology", "Medium", "In Progress", -9, -1, 0.9, -1, ""),
]


def main() -> None:
    today = date.today()
    wb = Workbook()
    ws = wb.active
    ws.title = "Tracker"

    thin = Side(style="thin", color="D0D5DD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="0B2545")

    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 28

    for i, row in enumerate(SAMPLE, start=2):
        ini, action, owner, func, prio, status, s_off, d_off, pct, u_off, remarks = row
        values = [
            f"A-{i - 1:03d}", ini, action, owner, func, prio, status,
            today + timedelta(days=s_off), today + timedelta(days=d_off), pct,
            today + timedelta(days=u_off), remarks,
        ]
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font = Font(name=FONT, size=10, color="0000FF")  # blue = typed input
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=c in (2, 3, 12))
        for c in (8, 9, 11):
            ws.cell(row=i, column=c).number_format = "dd-mmm-yyyy"
        ws.cell(row=i, column=10).number_format = "0%"

        # Formula columns (black = formula)
        ws.cell(row=i, column=13, value=f'=IF(G{i}="Completed","",I{i}-TODAY())')
        ws.cell(
            row=i, column=14,
            value=(
                f'=IF(G{i}="Completed","Done",IF(I{i}<TODAY(),"Red",'
                f'IF(OR(G{i}="Blocked",AND(I{i}-TODAY()<=7,J{i}<0.5)),"Amber","Green")))'
            ),
        )
        for c in (13, 14):
            cell = ws.cell(row=i, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=i, column=13).number_format = "0;-0;0"

    last = len(SAMPLE) + 1
    # Leave room for ~200 rows of dropdowns/formatting so new rows behave the same
    max_row = 200
    for col, options in (("E", FUNCTIONS), ("F", PRIORITIES), ("G", STATUSES)):
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{max_row}")

    rag_rng = f"N2:N{max_row}"
    for text, fill, font in (
        ("Red", "F8D7DA", "9C1C28"), ("Amber", "FFF0C2", "8A5A00"),
        ("Green", "D4EDDA", "1E6B3A"), ("Done", "E7EAEE", "475467"),
    ):
        ws.conditional_formatting.add(
            rag_rng,
            CellIsRule(operator="equal", formula=[f'"{text}"'],
                       fill=PatternFill("solid", bgColor=fill),
                       font=Font(name=FONT, bold=True, color=font)),
        )

    widths = [8, 28, 44, 16, 14, 10, 13, 13, 13, 11, 13, 34, 11, 9]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{last}"

    # ---- Instructions sheet ----
    ins = wb.create_sheet("Instructions")
    lines = [
        ("How to use this tracker", True),
        ("", False),
        ("1. The rows on the Tracker sheet are SAMPLE data. Replace them with your own actions.", False),
        ("2. Type only in columns A to L (blue text). Columns M and N are formulas - do not overwrite.", False),
        ("3. Status must be one of: Not Started, In Progress, Blocked, Completed (dropdown provided).", False),
        ("4. % Complete is entered as a percentage (e.g. 60%). Dates use dd-mmm-yyyy.", False),
        ("5. Run:  python report.py --input data/action_tracker.xlsx   to generate the weekly deck and follow-up emails.", False),
        ("", False),
        ("RAG rules (same logic in Excel and in report.py)", True),
        ("Done  : Status = Completed", False),
        ("Red   : not completed and Due Date is before today", False),
        ("Amber : Status = Blocked, OR due within 7 days with under 50% complete", False),
        ("Green : everything else", False),
    ]
    for r, (text, bold) in enumerate(lines, 1):
        cell = ins.cell(row=r, column=1, value=text)
        cell.font = Font(name=FONT, size=11, bold=bold)
    ins.column_dimensions["A"].width = 110

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Created {OUT} with {len(SAMPLE)} sample actions")


if __name__ == "__main__":
    main()

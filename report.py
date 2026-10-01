#!/usr/bin/env python3
"""Action Tracker -> automated weekly governance report.

Reads the Excel action tracker, calculates status / RAG metrics, and produces
  1. a 5-slide PowerPoint management update, and
  2. a text file of follow-up email drafts grouped by action owner.

Usage:
    python report.py --input data/action_tracker.xlsx [--as-of 2026-10-01] [--out reports]
"""
import argparse
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

REQUIRED = [
    "ID", "Initiative", "Action Item", "Owner", "Function", "Priority",
    "Status", "Due Date", "% Complete",
]
STATUSES = ["Not Started", "In Progress", "Blocked", "Completed"]
PRIORITY_RANK = {"High": 0, "Medium": 1, "Low": 2}
DUE_SOON_DAYS = 7

NAVY, BLUE, INK, GREY = "0B2545", "1F5FBF", "1D2939", "667085"
TINT, WHITE = "F2F5FA", "FFFFFF"
RAG_COLORS = {"Done": "98A2B3", "Green": "2E8B57", "Amber": "E0A100", "Red": "C0392B"}
FONT = "Calibri"


# ----------------------------------------------------------------- data layer
def load_tracker(path: Path, as_of: pd.Timestamp) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name="Tracker")
    missing = [c for c in REQUIRED if c not in df.columns]
    if missing:
        raise SystemExit(f"Tracker is missing required column(s): {', '.join(missing)}")

    df = df.dropna(subset=["ID"]).copy()
    df["Status"] = df["Status"].astype(str).str.strip().str.title()
    df["Priority"] = df["Priority"].astype(str).str.strip().str.title()
    df["Due Date"] = pd.to_datetime(df["Due Date"], errors="coerce")
    df["% Complete"] = pd.to_numeric(df["% Complete"], errors="coerce").fillna(0)
    if df["% Complete"].max() > 1:  # someone typed 60 instead of 60%
        df["% Complete"] = df["% Complete"] / 100

    # Data-quality checks: flag instead of silently failing
    problems = []
    for _, r in df.iterrows():
        if r["Status"] not in STATUSES:
            problems.append(f"{r['ID']}: unrecognised status '{r['Status']}'")
        if pd.isna(r["Due Date"]):
            problems.append(f"{r['ID']}: missing/invalid due date")
        if pd.isna(r["Owner"]) or not str(r["Owner"]).strip():
            problems.append(f"{r['ID']}: no owner assigned")
    if problems:
        print("DATA QUALITY WARNINGS (these rows are excluded from date-based metrics where relevant):")
        for p in problems:
            print("  -", p)

    df = df[df["Status"].isin(STATUSES)].copy()
    df["Owner"] = df["Owner"].fillna("Unassigned")
    df["Days to Due"] = (df["Due Date"] - as_of).dt.days
    df["Is Open"] = df["Status"] != "Completed"
    df["RAG"] = df.apply(_rag, axis=1)
    return df.reset_index(drop=True)


def _rag(r) -> str:
    """Same rules as the Excel formulas in the tracker."""
    if r["Status"] == "Completed":
        return "Done"
    if pd.notna(r["Days to Due"]) and r["Days to Due"] < 0:
        return "Red"
    if r["Status"] == "Blocked" or (
        pd.notna(r["Days to Due"]) and r["Days to Due"] <= DUE_SOON_DAYS and r["% Complete"] < 0.5
    ):
        return "Amber"
    return "Green"


def compute_metrics(df: pd.DataFrame) -> dict:
    total = len(df)
    done = int((df["Status"] == "Completed").sum())
    overdue = df[df["Is Open"] & (df["Days to Due"] < 0)]
    blocked = df[df["Status"] == "Blocked"]
    due_soon = df[df["Is Open"] & df["Days to Due"].between(0, DUE_SOON_DAYS)]
    attention = df[df["Is Open"] & ((df["Days to Due"] < 0) | (df["Status"] == "Blocked"))].copy()
    attention["_p"] = attention["Priority"].map(PRIORITY_RANK).fillna(3)
    attention = attention.sort_values(["_p", "Days to Due"])
    top_function = overdue["Function"].value_counts()
    return dict(
        total=total, done=done, done_pct=done / total if total else 0,
        in_progress=int((df["Status"] == "In Progress").sum()),
        not_started=int((df["Status"] == "Not Started").sum()),
        blocked=blocked, overdue=overdue, due_soon=due_soon.sort_values("Days to Due"),
        attention=attention,
        top_function=(top_function.index[0], int(top_function.iloc[0])) if len(top_function) else None,
        avg_open_progress=df.loc[df["Is Open"], "% Complete"].mean() if df["Is Open"].any() else 0,
    )


# ------------------------------------------------------------ pptx helpers
def rgb(hex_):
    return RGBColor.from_string(hex_)


def add_text(slide, text, x, y, w, h, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name, run.font.size, run.font.bold = FONT, Pt(size), bold
    run.font.color.rgb = rgb(color)
    return box


def new_slide(prs, title, as_of_label):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    add_text(slide, title, 0.6, 0.45, 12.1, 0.9, size=26, bold=True, color=NAVY,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, f"Source: Action Tracker, as of {as_of_label}", 0.6, 7.0, 8, 0.3,
             size=10, color=GREY)
    return slide


def add_table(slide, frame, columns, widths, x, y, row_h=0.46, rag_col=None):
    rows = len(frame) + 1
    shape = slide.shapes.add_table(rows, len(columns), Inches(x), Inches(y),
                                   Inches(sum(widths)), Inches(row_h * rows))
    tbl = shape.table
    for i, w in enumerate(widths):
        tbl.columns[i].width = Inches(w)
    for r in range(rows):
        tbl.rows[r].height = Inches(row_h)

    def fill_cell(cell, text, bold=False, color=INK, bg=WHITE, align=PP_ALIGN.LEFT):
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(bg)
        cell.margin_left = cell.margin_right = Inches(0.08)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = str(text)
        run.font.name, run.font.size, run.font.bold = FONT, Pt(12), bold
        run.font.color.rgb = rgb(color)

    for c, name in enumerate(columns):
        fill_cell(tbl.cell(0, c), name, bold=True, color=WHITE, bg=NAVY)
    for r, (_, row) in enumerate(frame.iterrows(), start=1):
        bg = WHITE if r % 2 else TINT
        for c, name in enumerate(columns):
            val = row[name]
            color, bold = INK, False
            if rag_col and name == rag_col:
                color, bold = RAG_COLORS.get(row["RAG"], INK), True
            fill_cell(tbl.cell(r, c), val, bold=bold, color=color, bg=bg)
    return shape


def shorten(text, n=52):
    text = str(text)
    return text if len(text) <= n else text[: n - 1].rstrip() + "…"


# ------------------------------------------------------------------ slides
def build_deck(df, m, as_of, out_path):
    label = as_of.strftime("%d %b %Y")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    # 1. Title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(NAVY)
    add_text(s, "Weekly Action Tracker Review", 0.9, 2.5, 11.5, 1.2, size=44, bold=True,
             color=WHITE, anchor=MSO_ANCHOR.BOTTOM)
    add_text(s, f"Strategic initiatives and governance actions  |  Week ending {label}",
             0.9, 3.9, 11.5, 0.6, size=20, color="C9D6EE")

    # 2. Summary
    headline = f"{m['done_pct']:.0%} of actions are complete; {len(m['overdue'])} are overdue"
    s = new_slide(prs, headline, label)
    cards = [
        ("Total actions", m["total"], NAVY), ("Completed", m["done"], RAG_COLORS["Green"]),
        ("In progress", m["in_progress"], BLUE), ("Blocked", len(m["blocked"]), RAG_COLORS["Amber"]),
        ("Overdue", len(m["overdue"]), RAG_COLORS["Red"]),
    ]
    card_w, gap, x0 = 2.25, 0.2, 0.6
    for i, (lab, val, col) in enumerate(cards):
        x = x0 + i * (card_w + gap)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.7),
                                  Inches(card_w), Inches(1.7))
        card.adjustments[0] = 0.08
        card.fill.solid()
        card.fill.fore_color.rgb = rgb(TINT)
        card.line.fill.background()
        card.shadow.inherit = False
        add_text(s, str(val), x, 1.85, card_w, 0.9, size=44, bold=True, color=col,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, lab, x, 2.8, card_w, 0.4, size=14, color=GREY, align=PP_ALIGN.CENTER)

    msgs = []
    if m["top_function"]:
        fn, n = m["top_function"]
        msgs.append(f"Overdue actions are concentrated in {fn} ({n} of {len(m['overdue'])}); "
                    f"consider a focused follow-up.")
    if len(m["blocked"]):
        msgs.append(f"{len(m['blocked'])} action(s) are blocked and need escalation or a management decision.")
    msgs.append(f"{len(m['due_soon'])} action(s) fall due in the next {DUE_SOON_DAYS} days.")
    msgs.append(f"Open actions are on average {m['avg_open_progress']:.0%} complete.")
    box = s.shapes.add_textbox(Inches(0.6), Inches(3.9), Inches(12.1), Inches(2.8))
    tf = box.text_frame
    tf.word_wrap = True
    add_text(s, "Key messages", 0.6, 3.75, 6, 0.4, size=18, bold=True, color=NAVY)
    box.top = Inches(4.25)
    for i, msg in enumerate(msgs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = "•  " + msg
        run.font.name, run.font.size = FONT, Pt(17)
        run.font.color.rgb = rgb(INK)

    # 3. Charts
    s = new_slide(prs, "Progress by status and by function", label)
    cd = CategoryChartData()
    cd.categories = STATUSES
    cd.add_series("Actions", [int((df["Status"] == st).sum()) for st in STATUSES])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.6),
                            Inches(5.2), Inches(5.2), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = "Actions by status"
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.name = FONT
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.category_axis.tick_labels.font.size = Pt(12)
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(14)
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    ser = plot.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = rgb(BLUE)

    funcs = sorted(df["Function"].dropna().unique())
    cd2 = CategoryChartData()
    cd2.categories = funcs
    for rag in ["Done", "Green", "Amber", "Red"]:
        cd2.add_series(rag if rag != "Green" else "On track",
                       [int(((df["Function"] == f) & (df["RAG"] == rag)).sum()) for f in funcs])
    gf2 = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(6.2), Inches(1.6),
                             Inches(6.5), Inches(5.2), cd2)
    ch2 = gf2.chart
    ch2.has_title = True
    ch2.chart_title.text_frame.text = "Actions by function and RAG status"
    ch2.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
    ch2.chart_title.text_frame.paragraphs[0].runs[0].font.name = FONT
    ch2.has_legend = True
    ch2.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch2.legend.include_in_layout = False
    ch2.legend.font.size = Pt(12)
    ch2.value_axis.visible = False
    ch2.value_axis.has_major_gridlines = False
    ch2.category_axis.tick_labels.font.size = Pt(12)
    ch2.category_axis.reverse_order = True
    p2 = ch2.plots[0]
    p2.gap_width = 50
    p2.has_data_labels = True
    p2.data_labels.number_format = "0;;"  # hide zeros
    p2.data_labels.number_format_is_linked = False
    p2.data_labels.position = XL_LABEL_POSITION.CENTER
    p2.data_labels.font.size = Pt(12)
    p2.data_labels.font.color.rgb = rgb(WHITE)
    for series, rag in zip(p2.series, ["Done", "Green", "Amber", "Red"]):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = rgb(RAG_COLORS[rag])

    # 4. Needs attention
    att = m["attention"]
    s = new_slide(prs, f"{len(m['overdue'])} overdue and {len(m['blocked'])} blocked actions need attention", label)
    if att.empty:
        add_text(s, "No overdue or blocked actions this week.", 0.6, 2.0, 12, 0.6, size=20)
    else:
        view = att.head(8).copy()
        view["Action Item"] = view["Action Item"].map(lambda t: shorten(t, 48))
        view["Due"] = view["Due Date"].dt.strftime("%d %b")
        view["Issue"] = view.apply(
            lambda r: f"Overdue {abs(int(r['Days to Due']))}d" if r["Days to Due"] < 0 else "Blocked", axis=1)
        view.loc[(view["Status"] == "Blocked") & (view["Days to Due"] < 0), "Issue"] = (
            view["Issue"] + " + blocked")
        cols = ["ID", "Action Item", "Owner", "Function", "Priority", "Due", "Issue"]
        add_table(s, view, cols, [0.9, 3.9, 1.7, 1.5, 1.0, 0.9, 2.2], 0.6, 1.6, rag_col="Issue")
        if len(att) > 8:
            add_text(s, f"+ {len(att) - 8} more in the tracker", 0.6, 6.6, 6, 0.3, size=12, color=GREY)

    # 5. Next 7 days
    ds = m["due_soon"]
    s = new_slide(prs, f"{len(ds)} actions fall due in the next {DUE_SOON_DAYS} days: follow-ups required", label)
    if ds.empty:
        add_text(s, f"Nothing due in the next {DUE_SOON_DAYS} days.", 0.6, 2.0, 12, 0.6, size=20)
    else:
        view = ds.head(8).copy()
        view["Action Item"] = view["Action Item"].map(lambda t: shorten(t, 48))
        view["Due"] = view["Due Date"].dt.strftime("%d %b")
        view["Progress"] = view["% Complete"].map(lambda v: f"{v:.0%}")
        cols = ["ID", "Action Item", "Owner", "Function", "Due", "Progress", "Status"]
        add_table(s, view, cols, [0.9, 4.5, 1.7, 1.5, 0.9, 1.1, 1.5], 0.6, 1.6, rag_col=None)

    prs.save(out_path)


# ---------------------------------------------------------- follow-up drafts
def write_followups(df, m, as_of, out_path):
    label = as_of.strftime("%d %b %Y")
    flagged = df[df["Is Open"] & (
        (df["Days to Due"] < 0) | (df["Status"] == "Blocked") | df["Days to Due"].between(0, DUE_SOON_DAYS)
    )]
    lines = [f"FOLLOW-UP EMAIL DRAFTS  |  as of {label}", "=" * 70, ""]
    for owner, grp in flagged.groupby("Owner"):
        grp = grp.sort_values("Days to Due")
        lines += [f"To: {owner}", f"Subject: Action tracker follow-up - {len(grp)} item(s) for your attention", "",
                  f"Hi {str(owner).split()[0]},", "",
                  "Please can you share a brief update on the following actions:", ""]
        for _, r in grp.iterrows():
            if r["Days to Due"] < 0:
                flag = f"OVERDUE by {abs(int(r['Days to Due']))} day(s)"
            elif r["Status"] == "Blocked":
                flag = "BLOCKED"
            else:
                flag = f"due in {int(r['Days to Due'])} day(s)"
            note = f" - {r['Remarks']}" if "Remarks" in df.columns and pd.notna(r.get("Remarks")) else ""
            lines.append(f"  - [{r['ID']}] {r['Action Item']} ({flag}, {r['% Complete']:.0%} complete){note}")
        lines += ["", "If you are blocked or the date needs to move, let me know so it can be escalated.",
                  "", "Thanks,", "COO Office", "", "-" * 70, ""]
    out_path.write_text("\n".join(lines), encoding="utf-8")
    return flagged["Owner"].nunique()


# --------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="Generate weekly report from the action tracker.")
    ap.add_argument("--input", default="data/action_tracker.xlsx", type=Path)
    ap.add_argument("--out", default="reports", type=Path)
    ap.add_argument("--as-of", default=None, help="Report date, YYYY-MM-DD (default: today)")
    args = ap.parse_args()

    as_of = pd.Timestamp(datetime.strptime(args.as_of, "%Y-%m-%d").date() if args.as_of else date.today())
    args.out.mkdir(parents=True, exist_ok=True)

    df = load_tracker(args.input, as_of)
    m = compute_metrics(df)

    stamp = as_of.strftime("%Y-%m-%d")
    deck = args.out / f"Weekly_Action_Report_{stamp}.pptx"
    mails = args.out / f"Follow_Up_Emails_{stamp}.txt"
    build_deck(df, m, as_of, deck)
    n_owners = write_followups(df, m, as_of, mails)

    print(f"\nActions: {m['total']} | Completed: {m['done']} ({m['done_pct']:.0%}) | "
          f"Overdue: {len(m['overdue'])} | Blocked: {len(m['blocked'])} | Due soon: {len(m['due_soon'])}")
    print(f"Deck:      {deck}")
    print(f"Follow-ups: {mails} ({n_owners} owners)")


if __name__ == "__main__":
    main()
